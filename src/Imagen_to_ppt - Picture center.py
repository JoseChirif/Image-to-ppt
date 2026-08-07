import os
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from PIL import Image

# go to the parent directory if you are running this script directly (uncomment the following lines)
# import sys
# sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from config.config import pictures_extensions
from functions.functions import working_directory, show_options, show_message, check_presentation_exists, run_task_with_progress



def create_presentation():
    """
    Creates a PowerPoint presentation from images in the working directory.
    """
    # Variables
    ppt_name = "Presentation - Pictures center ppt"

    # Get the working directory and construct the output file path
    image_folder = working_directory()
    output_pptx = os.path.join(
        image_folder,
        ppt_name + ".pptx",
    )

    # Check if the file exists and handle replacement decision
    if not check_presentation_exists(output_pptx):
        show_message(
            "Operation canceled",
            "The presentation was not created.",
        )
        return

    # Get the list of images in the folder
    images = [
        image
        for image in os.listdir(image_folder)
        if image.lower().endswith(pictures_extensions)
    ]

    # Get option nr
    option_nr = show_options(
        "Do you want to include \n the file name's in the slide?",
        "Yes",
        "No",
    )

    # If option_nr is None, cancel the operation
    if option_nr is None:
        show_message(
            "Operation canceled",
            "No presentation was created.",
        )
        return

    # -------------------------------------------------------------------------
    # This function runs in the worker thread.
    # -------------------------------------------------------------------------
    def build_presentation(report_progress):
        # Create a blank presentation with standard dimensions
        prs = Presentation()
        prs.slide_width = Cm(33.867)
        prs.slide_height = Cm(19.05)

        for image_number, image in enumerate(
            images,
            start=1,
        ):
            # Full path of the image
            image_path = os.path.join(
                image_folder,
                image,
            )

            # Get image dimensions
            with Image.open(image_path) as img:
                img_width, img_height = img.size

            # Set the fixed height for the image
            new_height = Cm(17.43)
            new_width = int(
                (img_width / img_height) * new_height
            )

            if new_width > prs.slide_width:
                new_width = prs.slide_width
                new_height = int(
                    (img_height / img_width) * new_width
                )

            # Calculate positions to center the image
            left = int(
                (prs.slide_width - new_width) / 2
            )

            top = int(
                (prs.slide_height - new_height) / 2
            )

            # Create a new slide
            slide_layout = prs.slide_layouts[6]
            slide = prs.slides.add_slide(slide_layout)

            # Add the centered image with a black border
            picture = slide.shapes.add_picture(
                image_path,
                left,
                top,
                width=new_width,
                height=new_height,
            )

            line = picture.line
            line.color.rgb = RGBColor(0, 0, 0)
            line.width = Pt(0.75)

            if option_nr == 1:
                # Get the file name without the extension
                file_name = os.path.splitext(image)[0]

                # Add the file name at the bottom of the slide
                text_left = Cm(0)
                text_top = prs.slide_height - Cm(1.5)
                text_width = prs.slide_width
                text_height = Cm(1)

                txBox = slide.shapes.add_textbox(
                    text_left,
                    text_top,
                    text_width,
                    text_height,
                )

                tf = txBox.text_frame
                tf.clear()

                paragraph = tf.add_paragraph()
                paragraph.text = file_name
                paragraph.font.size = Pt(10)
                paragraph.font.color.rgb = RGBColor(
                    0,
                    0,
                    0,
                )
                paragraph.font.name = "Aptos"
                paragraph.alignment = 2

            elif option_nr == 2:
                pass

            # Report progress only after the image was fully processed.
            report_progress(
                image_number,
                image,
            )

        # Inform the popup that the slides are complete
        # and the PowerPoint is being saved.
        report_progress(
            len(images),
            os.path.basename(output_pptx),
            "Saving presentation...",
        )

        prs.save(output_pptx)

    # -------------------------------------------------------------------------
    # Open the popup and execute the worker.
    # -------------------------------------------------------------------------
    run_task_with_progress(
        total_items=len(images),
        task=build_presentation,
        title="Creating presentation",
    )

    show_message(
        "Done",
        f"Presentation saved as {ppt_name + '.pptx'}",
    )



if __name__ == "__main__":
    create_presentation()
